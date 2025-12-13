import uuid
import logging
from typing import List, Dict, Any, Optional
from io import BytesIO

# This file:
# 1. Takes a PPTX or PDF slide deck 
# 2. Extracts structured content
# 3. Normalizes it 
# 4. Optionally OCRs images 
# 5. Converts everything into clean, searchable text chunks 
# 6. Stores them in a vector DB

# RAG DESIGN:
# We have two chunking strategies:
# 1. Slide chunks: 
# These are the chunks that are a single slide, we use this for precise retrieval
# 2. Window chunks: 
# These are the chunks that are a range of slides, we use this for broader topic retrieval


# These are the libraries that we use to load the slides
# Loaded them conditionally in case the someone running this doesn't have one of them installed
try:
    from pptx import Presentation
    from pptx.shapes.picture import Picture
    PPTX_AVAILABLE = True
except ImportError:
    PPTX_AVAILABLE = False

try:
    import fitz
    PYMuPDF_AVAILABLE = True
except ImportError:
    PYMuPDF_AVAILABLE = False

from .ocr_adapter import OCRAdapter, get_ocr_adapter, should_ocr_slide
from .slide_normalization import normalize_bullet_characters, denoise_ocr_text
from .slide_models import TextBlock, ImageBlock, SlideData

logger = logging.getLogger(__name__)


# (Window Size - Window Stride) / Window Size = Overlap Ratio
# Example: (10 - 7) / 10 = 0.3 = 30% overlap aka 3 slides overlap
# window 1 
# window 1: slides 1-10
# window 2: slides 8-17
# etc...
WINDOW_SIZE = 10 # number of slides per window
WINDOW_STRIDE = 7 # number of slides to stride between windows


# Loads PPTX file and extracts text/images from each slide
# Returns list of SlideData objects
def load_pptx(file_bytes: bytes, deck_id: str) -> List[SlideData]:
    if not PPTX_AVAILABLE:
        raise ImportError("python-pptx is required. Install with: pip install python-pptx")
    
    try:
        prs = Presentation(BytesIO(file_bytes))
        slides_data = []
        
        for idx, slide in enumerate(prs.slides, start=1):
            # 1. Get slide dimensions (convert EMU to inches)
            slide_width = prs.slide_width / 914400
            slide_height = prs.slide_height / 914400
            
            text_blocks = []
            image_blocks = []
            
            # 2. Extract text from shapes
            for shape in slide.shapes:
                if shape.has_text_frame:
                    text_frame = shape.text_frame
                    full_text = ""
                    
                    for paragraph in text_frame.paragraphs:
                        para_text = paragraph.text.strip()
                        if para_text:
                            full_text += para_text + "\n"
                    
                    if full_text.strip():
                        # 3. Get bounding box (EMU to inches)
                        left_pt = shape.left / 914400
                        top_pt = shape.top / 914400
                        width_pt = shape.width / 914400
                        height_pt = shape.height / 914400
                        
                        # 4. Try to get font size
                        font_size = None
                        if text_frame.paragraphs:
                            first_para = text_frame.paragraphs[0]
                            if first_para.runs:
                                font_size = first_para.runs[0].font.size
                                if font_size:
                                    font_size = font_size / 12700
                        
                        text_blocks.append(TextBlock(
                            text=full_text.strip(),
                            bbox=[left_pt, top_pt, width_pt, height_pt],
                            font_size=font_size
                        ))
                
                # 5. Extract images
                if isinstance(shape, Picture):
                    try:
                        image_bytes = shape.image.blob
                        left_pt = shape.left / 914400
                        top_pt = shape.top / 914400
                        width_pt = shape.width / 914400
                        height_pt = shape.height / 914400
                        
                        image_blocks.append(ImageBlock(
                            image_bytes=image_bytes,
                            bbox=[left_pt, top_pt, width_pt, height_pt]
                        ))
                    except Exception as e:
                        logger.warning(f"Failed to extract image from slide {idx}: {e}")
            
            # 6. Extract notes if available
            notes_text = None
            if slide.has_notes_slide:
                notes_slide = slide.notes_slide
                if notes_slide.notes_text_frame:
                    notes_text = notes_slide.notes_text_frame.text.strip()
            
            slides_data.append(SlideData(
                deck_id=deck_id,
                slide_number=idx,
                width=slide_width,
                height=slide_height,
                text_blocks=text_blocks,
                image_blocks=image_blocks,
                notes_text=notes_text
            ))
        
        return slides_data
    
    except Exception as e:
        raise ValueError(f"Failed to parse PPTX file: {e}")


# Loads PDF and treats each page as a slide
# Extracts text blocks and images from each page
def load_pdf_slides(file_bytes: bytes, deck_id: str) -> List[SlideData]:
    if not PYMuPDF_AVAILABLE:
        raise ImportError("PyMuPDF (fitz) is required. Install with: pip install pymupdf")
    
    try:
        doc = fitz.open(stream=file_bytes, filetype="pdf")
        slides_data = []
        
        for page_num in range(len(doc)):
            page = doc[page_num]
            
            # 1. Get page dimensions (convert points to inches: 1 inch = 72 points)
            rect = page.rect
            page_width = rect.width / 72.0
            page_height = rect.height / 72.0
            
            text_blocks = []
            image_blocks = []
            
            # 2. Extract text blocks
            text_dict = page.get_text("dict")
            for block in text_dict.get("blocks", []):
                if block.get("type") == 0:
                    block_text = ""
                    bbox_points = block.get("bbox", [0, 0, 0, 0])
                    
                    for line in block.get("lines", []):
                        for span in line.get("spans", []):
                            span_text = span.get("text", "").strip()
                            if span_text:
                                block_text += span_text + " "
                    
                    block_text = block_text.strip()
                    if block_text:
                        # 3. Convert bbox from points to inches [x0, y0, x1, y1] -> [left, top, width, height]
                        left = bbox_points[0] / 72.0
                        top = bbox_points[1] / 72.0
                        width = (bbox_points[2] - bbox_points[0]) / 72.0
                        height = (bbox_points[3] - bbox_points[1]) / 72.0
                        bbox_inches = [left, top, width, height]
                        
                        # 4. Get font size if available
                        font_size = None
                        if block.get("lines"):
                            first_line = block["lines"][0]
                            if first_line.get("spans"):
                                font_size = first_line["spans"][0].get("size")
                        
                        text_blocks.append(TextBlock(
                            text=block_text,
                            bbox=bbox_inches,
                            font_size=font_size
                        ))
            
            # 5. Extract images
            image_list = page.get_images(full=True)
            for img_idx, img in enumerate(image_list):
                try:
                    xref = img[0]
                    base_image = doc.extract_image(xref)
                    image_bytes = base_image["image"]
                    # Convert image bbox to inches (using full page for now)
                    bbox = [0, 0, page_width, page_height]
                    
                    image_blocks.append(ImageBlock(
                        image_bytes=image_bytes,
                        bbox=bbox
                    ))
                except Exception as e:
                    logger.warning(f"Failed to extract image {img_idx} from page {page_num + 1}: {e}")
            
            slides_data.append(SlideData(
                deck_id=deck_id,
                slide_number=page_num + 1,
                width=page_width,
                height=page_height,
                text_blocks=text_blocks,
                image_blocks=image_blocks,
                notes_text=None
            ))
        
        doc.close()
        return slides_data
    
    except Exception as e:
        raise ValueError(f"Failed to parse PDF file: {e}")





# Runs OCR on slide images and returns text blocks
# Returns empty list if OCR unavailable
def apply_ocr_to_slide(slide: SlideData, ocr_adapter: Optional[OCRAdapter]) -> List[TextBlock]:
    if ocr_adapter is None:
        return []
    
    ocr_blocks = []
    
    try:
        # 1. OCR each image block
        for img_block in slide.image_blocks:
            if img_block.image_bytes:
                try:
                    ocr_text = ocr_adapter.extract_text(img_block.image_bytes)
                    if ocr_text and ocr_text.strip():
                        ocr_blocks.append(TextBlock(
                            text=ocr_text.strip(),
                            bbox=img_block.bbox.copy(),
                            font_size=None
                        ))
                except Exception as e:
                    logger.warning(f"OCR failed for image on slide {slide.slide_number}: {e}")
    except Exception as e:
        logger.error(f"OCR processing failed for slide {slide.slide_number}: {e}")
    
    return ocr_blocks


# Finds the slide title from text blocks
# Uses heuristic: largest font size near top
def detect_slide_title(text_blocks: List[TextBlock]) -> Optional[str]:
    if not text_blocks:
        return None
    
    # 1. Sort by vertical position, then font size
    sorted_blocks = sorted(
        text_blocks,
        key=lambda b: (b.bbox[1], -(b.font_size or 0))
    )
    
    # 2. Take first line of top block as title
    if sorted_blocks:
        title = sorted_blocks[0].text.split('\n')[0].strip()
        if len(title) < 200:
            return title
    
    return None


# Builds canonical text representation for a slide
# Format: header, title, bullets, notes, OCR text
"""

[Deck: X] [Slide: Y]

Title: ...

Bullets:
- ...
- ...

Notes:
- ...

Image Text (OCR):
- ...

"""

def build_canonical_slide_text(slide: SlideData) -> str:
    parts = []
    
    # 1. Add header
    parts.append(f"[Deck: {slide.deck_id}] [Slide: {slide.slide_number}]")
    parts.append("")
    
    # 2. Add title if found
    title = detect_slide_title(slide.text_blocks)
    if title:
        parts.append(f"Title: {title}")
        parts.append("")
    
    # 3. Collect native text blocks (skip title if already added)
    native_text_lines = []
    title_included = False
    
    for block in slide.text_blocks:
        block_text = block.text.strip()
        if not block_text:
            continue
        
        if title and not title_included:
            first_line = block_text.split('\n')[0].strip()
            if first_line == title or first_line in title:
                title_included = True
                remaining = block_text[len(first_line):].strip()
                if remaining:
                    native_text_lines.append(remaining)
                continue
        
        native_text_lines.append(block_text)
    
    # 4. Format as bullets
    if native_text_lines:
        parts.append("Bullets:")
        for line in native_text_lines:
            normalized_line = normalize_bullet_characters(line)
            for subline in normalized_line.split('\n'):
                subline_stripped = subline.strip()
                if subline_stripped:
                    if not subline_stripped.startswith('- '):
                        subline_stripped = '- ' + subline_stripped
                    parts.append(subline_stripped)
        parts.append("")
    
    # 5. Add notes if available
    if slide.notes_text and slide.notes_text.strip():
        parts.append("Notes:")
        notes_normalized = normalize_bullet_characters(slide.notes_text.strip())
        for line in notes_normalized.split('\n'):
            if line.strip():
                parts.append(f"- {line.strip()}")
        parts.append("")
    
    # 6. Add OCR text if available
    if slide.ocr_text_blocks:
        parts.append("Image Text (OCR):")
        for block in slide.ocr_text_blocks:
            ocr_text = denoise_ocr_text(block.text)
            if ocr_text.strip():
                for line in ocr_text.split('\n'):
                    if line.strip():
                        parts.append(f"- {line.strip()}")
        parts.append("")
    
    return '\n'.join(parts).strip()


# Creates one chunk per slide with canonical text
# Used for precise retrieval
def build_slide_chunks(slides: List[SlideData], deck_id: str) -> List[Dict[str, Any]]:
    chunks = []
    
    for slide in slides:
        # 1. Build canonical text
        canonical_text = build_canonical_slide_text(slide)
        if not canonical_text.strip():
            continue
        
        # 2. Detect title for metadata
        slide_title = None
        if slide.text_blocks:
            sorted_blocks = sorted(
                slide.text_blocks,
                key=lambda b: (b.bbox[1], -(b.font_size or 0))
            )
            if sorted_blocks:
                slide_title = sorted_blocks[0].text.split('\n')[0].strip()
        
        # 3. Create chunk
        chunks.append({
            "text": canonical_text,
            "metadata": {
                "chunk_type": "slide",
                "deck_id": deck_id,
                "slide_number": slide.slide_number,
                "slide_title": slide_title or "",
                "has_ocr": len(slide.ocr_text_blocks) > 0,
                "source": deck_id,
            }
        })
    
    return chunks


# Creates window chunks with 10 slides per chunk, 30% overlap
# Used for broader topic retrieval
def build_window_chunks(slides: List[SlideData], deck_id: str, 
                       window_size: int = WINDOW_SIZE, stride: int = WINDOW_STRIDE) -> List[Dict[str, Any]]:
    chunks = []
    start_idx = 0
    
    while start_idx < len(slides):
        # 1. Get window of slides
        end_idx = min(start_idx + window_size, len(slides))
        window_slides = slides[start_idx:end_idx]
        
        # 2. Build canonical text for each slide
        slide_texts = []
        for slide in window_slides:
            canonical_text = build_canonical_slide_text(slide)
            if canonical_text.strip():
                slide_texts.append(canonical_text)
        
        # 3. Combine and create chunk
        if slide_texts:
            combined_text = "\n\n---\n\n".join(slide_texts)
            chunks.append({
                "text": combined_text,
                "metadata": {
                    "chunk_type": "window",
                    "deck_id": deck_id,
                    "start_slide": window_slides[0].slide_number,
                    "end_slide": window_slides[-1].slide_number,
                    "source": deck_id,
                }
            })
        
        start_idx += stride
    
    return chunks


# Main function to ingest PDF deck
# Loads slides, applies OCR, chunks, and stores in ChromaDB
def ingest_pdf_deck(file_bytes: bytes, deck_id: Optional[str] = None,
                   ocr_adapter: Optional[OCRAdapter] = None, store_in_db: bool = True) -> Dict[str, Any]:
    if deck_id is None:
        deck_id = f"pdf_deck_{uuid.uuid4().hex[:8]}"
    
    logger.info(f"Starting PDF deck ingestion: {deck_id}")
    
    # 1. Load slides from PDF
    slides = load_pdf_slides(file_bytes, deck_id)
    logger.info(f"Loaded {len(slides)} slides from PDF")
    
    # 2. Apply OCR where needed
    ocr_stats = {"slides_ocred": 0, "total_ocr_chars": 0}
    for slide in slides:
        if should_ocr_slide(slide):
            ocr_blocks = apply_ocr_to_slide(slide, ocr_adapter)
            if ocr_blocks:
                slide.ocr_text_blocks = ocr_blocks
                ocr_stats["slides_ocred"] += 1
                ocr_stats["total_ocr_chars"] += sum(len(block.text) for block in ocr_blocks)
    
    # 3. Build chunks
    slide_chunks = build_slide_chunks(slides, deck_id)
    window_chunks = build_window_chunks(slides, deck_id)
    logger.info(f"Created {len(slide_chunks)} slide chunks and {len(window_chunks)} window chunks")
    
    # 4. Store in ChromaDB
    stored_count = 0
    if store_in_db:
        try:
            import chromadb
            from config import CHROMA_PATH, COLLECTION_NAME
            from .embeddings import LocalEmbeddingFunction
            
            client = chromadb.PersistentClient(path=CHROMA_PATH)
            collection = client.get_or_create_collection(
                name=COLLECTION_NAME,
                embedding_function=LocalEmbeddingFunction(),
            )
            
            all_chunks = slide_chunks + window_chunks
            ids, docs, metas = [], [], []
            
            for chunk in all_chunks:
                chunk_type = chunk["metadata"]["chunk_type"]
                if chunk_type == "slide":
                    chunk_id = f"{deck_id}::s{chunk['metadata']['slide_number']}::slide"
                else:
                    chunk_id = f"{deck_id}::w{chunk['metadata']['start_slide']}-{chunk['metadata']['end_slide']}::window"
                
                ids.append(chunk_id)
                docs.append(chunk["text"])
                metas.append(chunk["metadata"])
            
            if ids:
                collection.upsert(ids=ids, documents=docs, metadatas=metas)
                stored_count = len(ids)
                logger.info(f"Stored {stored_count} chunks in ChromaDB")
        
        except Exception as e:
            logger.error(f"Failed to store chunks in ChromaDB: {e}")
            return {
                "deck_id": deck_id,
                "slides_processed": len(slides),
                "slide_chunks_count": len(slide_chunks),
                "window_chunks_count": len(window_chunks),
                "stored_count": 0,
                "ocr_stats": ocr_stats,
                "warnings": [f"Storage failed: {str(e)}"]
            }
    
    return {
        "deck_id": deck_id,
        "slides_processed": len(slides),
        "slide_chunks_count": len(slide_chunks),
        "window_chunks_count": len(window_chunks),
        "stored_count": stored_count,
        "ocr_stats": ocr_stats,
        "warnings": []
    }


# Main function to ingest PPTX deck
# Same as PDF but uses PPTX loader
def ingest_pptx_deck(file_bytes: bytes, deck_id: Optional[str] = None,
                    ocr_adapter: Optional[OCRAdapter] = None, store_in_db: bool = True) -> Dict[str, Any]:
    if deck_id is None:
        deck_id = f"pptx_deck_{uuid.uuid4().hex[:8]}"
    
    logger.info(f"Starting PPTX deck ingestion: {deck_id}")
    
    # 1. Load slides
    slides = load_pptx(file_bytes, deck_id)
    logger.info(f"Loaded {len(slides)} slides from PPTX")
    
    # 2. Apply OCR
    ocr_stats = {"slides_ocred": 0, "total_ocr_chars": 0}
    warnings = []
    
    if ocr_adapter is None:
        warnings.append("OCR adapter not available - image-heavy slides may have missing text")
    
    for slide in slides:
        if should_ocr_slide(slide):
            ocr_blocks = apply_ocr_to_slide(slide, ocr_adapter)
            if ocr_blocks:
                slide.ocr_text_blocks = ocr_blocks
                ocr_stats["slides_ocred"] += 1
                ocr_stats["total_ocr_chars"] += sum(len(block.text) for block in ocr_blocks)
    
    # 3. Build chunks
    slide_chunks = build_slide_chunks(slides, deck_id)
    window_chunks = build_window_chunks(slides, deck_id)
    logger.info(f"Created {len(slide_chunks)} slide chunks and {len(window_chunks)} window chunks")
    
    # 4. Store in ChromaDB
    stored_count = 0
    if store_in_db:
        try:
            import chromadb
            from config import CHROMA_PATH, COLLECTION_NAME
            from .embeddings import LocalEmbeddingFunction
            
            client = chromadb.PersistentClient(path=CHROMA_PATH)
            collection = client.get_or_create_collection(
                name=COLLECTION_NAME,
                embedding_function=LocalEmbeddingFunction(),
            )
            
            all_chunks = slide_chunks + window_chunks
            ids, docs, metas = [], [], []
            
            for chunk in all_chunks:
                chunk_type = chunk["metadata"]["chunk_type"]
                if chunk_type == "slide":
                    chunk_id = f"{deck_id}::s{chunk['metadata']['slide_number']}::slide"
                else:
                    chunk_id = f"{deck_id}::w{chunk['metadata']['start_slide']}-{chunk['metadata']['end_slide']}::window"
                
                ids.append(chunk_id)
                docs.append(chunk["text"])
                metas.append(chunk["metadata"])
            
            if ids:
                collection.upsert(ids=ids, documents=docs, metadatas=metas)
                stored_count = len(ids)
                logger.info(f"Stored {stored_count} chunks in ChromaDB")
        
        except Exception as e:
            logger.error(f"Failed to store chunks in ChromaDB: {e}")
            return {
                "deck_id": deck_id,
                "slides_processed": len(slides),
                "slide_chunks_count": len(slide_chunks),
                "window_chunks_count": len(window_chunks),
                "stored_count": 0,
                "ocr_stats": ocr_stats,
                "warnings": warnings + [f"Storage failed: {str(e)}"]
            }
    
    return {
        "deck_id": deck_id,
        "slides_processed": len(slides),
        "slide_chunks_count": len(slide_chunks),
        "window_chunks_count": len(window_chunks),
        "stored_count": stored_count,
        "ocr_stats": ocr_stats,
        "warnings": warnings
    }
